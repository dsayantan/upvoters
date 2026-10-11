#!/usr/bin/env python3
"""
Infer families from the voter workbooks of a UP assembly, booth by booth --
the rules of ~/VOTER_LIST_2026_XLSX/kmc_family.py, for Hindi rolls.

  ./up_family.py 186_Amethi_xlsx                 # -> FAMILY/186_Amethi_Families.xlsx
  ./up_family.py 186_Amethi_xlsx -o FAMILY -j 6

Rules (inside one booth; deleted voters are skipped)
  1. CHILD  : X's Father/Mother name matches a voter Y of the right gender who
              is at least MIN_PARENT_GAP years older        -> X child of Y
  2. SPOUSE : X's Husband name matches a male voter Y of plausible age
                                                            -> X wife of Y
  3. SIBLING: X and Y have the same Father (or same Mother) name -> siblings
  4. MOTHER : X's Husband name equals Y's Father name        -> X mother of Y
  Names match exactly after normalisation (spacing, honorifics such as श्री,
  स्व0, मो0, and the spelling folds of up_names.fold).  A looser consonant
  skeleton absorbs OCR noise, but only for voters in the same house or within
  NEAR serial numbers.  Rules 3 and 4 need the same house when both house
  numbers are known (49 and 49अ count as one), else NEAR serials -- two men of
  one common name in different houses are not merged through their children.
  Links are joined strongest first (union-find); a link is refused if the
  family it would make holds two women with the same husband name, or spans
  more than MAX_SPAN serial numbers.  Every group of 2+ is a family.

Output (FAMILY/)
  <AC>_<Name>_Families.xlsx   sheets Summary (per booth), Families, Members
  <AC>_<Name>_Families.json   {part: [[family no, [serial, ...], [links, ...]], ...]}
                              read by build_voters_local.py

These files hold voters' personal details: keep them on this machine.
"""

from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys
import unicodedata
from collections import defaultdict
from multiprocessing import Pool

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from up_names import fold

NEAR = 40               # serial-number window counted as "close by"
MIN_PARENT_GAP = 18     # parent must be this much older than child
MAX_SPAN = 100          # serial numbers inside one family span at most this
DELETED = {"E-Dead", "S-Shifted", "R-Repeated", "M-Missing", "Q-Ineligible"}

HONORIFICS = {"श्री", "श्रीमती", "स्व", "स्वर्गीय", "स्वo", "मो", "मोहम्मद", "मुहम्मद",
              "डा", "डॉ", "डॉक्टर", "कु", "late", "lt", "sri", "shri", "smt", "md", "mohd"}
DEVA_CONS = re.compile("[\u0915-\u0939\u0958-\u095f]")


def tokens(name):
    if not name:
        return []
    s = unicodedata.normalize("NFC", str(name)).replace("‌", "").replace("‍", "")
    s = re.sub(r"[^\wऀ-ॿ ]+", " ", s.lower())
    s = re.sub(r"[\d_०-९]+", " ", s)
    return [t for t in s.split() if t not in HONORIFICS]


def exact_key(name):
    return "".join(fold(t) or t for t in tokens(name))


def skeleton(name):
    s = exact_key(name)
    out = "".join(DEVA_CONS.findall(s))
    out = re.sub(r"(.)\1+", r"\1", out)
    return out if len(out) >= 3 else ""


def house_key(h):
    s = re.sub(r"\s+", "", str(h or "")).lower()
    return "" if s in ("", "0", "00", "000", "-") else s     # "00" is no house number


class DSU:
    """Union-find whose components remember their serial range and the husband
    names of their married women, so a union that breaks a rule is refused."""

    def __init__(self, voters):
        n = len(voters)
        self.p = list(range(n))
        self.lo = [v["sl"] for v in voters]
        self.hi = [v["sl"] for v in voters]
        self.hk = [set() for _ in range(n)]
        for i, v in enumerate(voters):
            if v["rtype"] == "husband" and v["rk"]:
                self.hk[i] = {("k", v["rk"])} | ({("s", v["rs"])} if v["rs"] else set())

    def find(self, a):
        while self.p[a] != a:
            self.p[a] = self.p[self.p[a]]
            a = self.p[a]
        return a

    def union(self, a, b):
        a, b = self.find(a), self.find(b)
        if a == b:
            return True
        los = [x for x in (self.lo[a], self.lo[b]) if x is not None]
        his = [x for x in (self.hi[a], self.hi[b]) if x is not None]
        if los and max(his) - min(los) > MAX_SPAN:
            return False
        if self.hk[a] & self.hk[b]:
            return False
        a, b = min(a, b), max(a, b)
        self.p[b] = a
        self.lo[a] = min(los) if los else None
        self.hi[a] = max(his) if his else None
        self.hk[a] |= self.hk[b]
        self.hk[b] = set()
        return True


def as_int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def read_part(path):
    """Every voter row of one booth workbook, deleted ones included (flagged)."""
    book = load_workbook(path, read_only=True)
    ws = book["Voters"] if "Voters" in book.sheetnames else book.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h else "" for h in next(rows)]
    ix = {h: i for i, h in enumerate(head)}
    voters = []
    for r in rows:
        g = lambda k: r[ix[k]] if k in ix and ix[k] < len(r) else None
        if not (g("Name") or g("Voter ID")):
            continue
        mark = (g("Deleted") or "").strip()
        voters.append({
            "sl": as_int(g("Serial No")), "epic": (g("Voter ID") or "").strip(),
            "name": (g("Name") or "").strip(), "rtype": (g("Relation Type") or "").strip().lower(),
            "rname": (g("Relation Name") or "").strip(), "house": str(g("House No") or "").strip(),
            "section": (g("Section") or "").strip(), "list": (g("List") or "").strip(),
            "age": as_int(g("Age")), "gender": (g("Gender") or "").strip(),
            "mark": mark, "deleted": mark in DELETED,
        })
    book.close()
    return voters


def infer_families(voters):
    """[[voter index, ...], ...] families of 2+, and links per voter index."""
    live = [i for i, v in enumerate(voters) if not v["deleted"]]
    vs = [voters[i] for i in live]
    for v in vs:
        v["k"], v["s"] = exact_key(v["name"]), skeleton(v["name"])
        v["rk"], v["rs"] = exact_key(v["rname"]), skeleton(v["rname"])
        v["h"] = house_key(v["house"])
    by_k, by_s = defaultdict(list), defaultdict(list)
    for i, v in enumerate(vs):
        if v["k"]:
            by_k[v["k"]].append(i)
        if v["s"]:
            by_s[v["s"]].append(i)

    def near(a, b):
        va, vb = vs[a], vs[b]
        if va["h"] and va["h"] == vb["h"]:
            return True
        return va["sl"] is not None and vb["sl"] is not None and abs(va["sl"] - vb["sl"]) <= NEAR

    def housemates(a, b):
        """Shared-guardian links need the same house when both houses are known
        (49 and 49अ count as one); without a house, NEAR serials will do."""
        ha, hb = vs[a]["h"], vs[b]["h"]
        if ha and hb:
            na, nb = re.match(r"\d*", ha).group(), re.match(r"\d*", hb).group()
            return ha == hb or (na != "" and na == nb)
        return near(a, b)

    def dist(a, b):
        va, vb = vs[a], vs[b]
        same = 0 if va["h"] and va["h"] == vb["h"] else 1
        d = abs(va["sl"] - vb["sl"]) if va["sl"] is not None and vb["sl"] is not None else 10**6
        return (same, d)

    def tag(i):
        return str(vs[i]["sl"])

    cands = []
    for x, v in enumerate(vs):
        rt = v["rtype"]
        if rt not in ("father", "mother", "husband") or not v["rk"]:
            continue
        want = "Female" if rt == "mother" else "Male"
        exact = [y for y in by_k.get(v["rk"], []) if y != x]
        loose = [y for y in by_s.get(v["rs"], []) if y != x and y not in exact and near(x, y)] \
            if v["rs"] else []

        def ok(y):
            w = vs[y]
            if w["gender"] and w["gender"] != want:
                return False
            if v["age"] is None or w["age"] is None:
                return True
            if rt == "husband":
                return v["age"] - 12 <= w["age"] <= v["age"] + 35
            return w["age"] - v["age"] >= MIN_PARENT_GAP

        for tier, ys in ((0, exact), (1, loose)):
            for y in ys:
                if not ok(y):
                    continue
                if rt == "husband":
                    cands.append(((tier,) + dist(x, y), x, y, ("wife", tag(y)), ("husband", tag(x)), x))
                else:
                    cands.append(((tier,) + dist(x, y), x, y, ("child", tag(y)),
                                  ("father" if rt == "father" else "mother", tag(x)), x))

    groups = defaultdict(list)
    for i, v in enumerate(vs):
        if v["rtype"] in ("father", "husband") and v["rk"]:
            groups[("M", v["rk"])].append(i)
        elif v["rtype"] == "mother" and v["rk"]:
            groups[("F", v["rk"])].append(i)
    for mem in groups.values():
        for ai in range(len(mem)):
            for bi in range(ai + 1, len(mem)):
                a, b = mem[ai], mem[bi]
                if not housemates(a, b):
                    continue
                ra, rb = vs[a]["rtype"], vs[b]["rtype"]
                if ra == rb == "husband":
                    continue
                if "husband" in (ra, rb):
                    m, c = (a, b) if ra == "husband" else (b, a)
                    if vs[m]["age"] and vs[c]["age"] and vs[m]["age"] - vs[c]["age"] < MIN_PARENT_GAP:
                        continue
                    cands.append(((2,) + dist(m, c), m, c, ("mother", tag(c)), ("child", tag(m)), None))
                else:
                    cands.append(((2,) + dist(a, b), a, b, ("sibling", tag(b)), ("sibling", tag(a)), None))

    cands.sort(key=lambda c: c[0])
    dsu = DSU(vs)
    explained = set()
    links = defaultdict(list)
    for _, a, b, la, lb, owner in cands:
        if owner is not None and owner in explained:
            continue
        if not dsu.union(a, b):
            continue
        if owner is not None:
            explained.add(owner)
        links[a].append(la)
        links[b].append(lb)

    fam = defaultdict(list)
    for i in range(len(vs)):
        fam[dsu.find(i)].append(i)
    out = []
    for mem in fam.values():
        if len(mem) < 2:
            continue
        mem.sort(key=lambda i: (-(vs[i]["age"] or 0), vs[i]["sl"] or 0))
        out.append(mem)
    out.sort(key=lambda m: min(vs[i]["sl"] or 0 for i in m))
    # back to indices into the full voter list
    return [[live[i] for i in m] for m in out], {live[i]: l for i, l in links.items()}


def part_no(path):
    return int(re.findall(r"\d+", os.path.basename(path))[-1])


def process(path):
    voters = read_part(path)
    fams, links = infer_families(voters)
    return part_no(path), voters, fams, links


LINK_HI = {"wife": "पत्नी", "husband": "पति", "child": "संतान", "father": "पिता",
           "mother": "माता", "sibling": "भाई/बहन"}


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[1])
    ap.add_argument("xlsx_dir", help="booth workbooks, e.g. 186_Amethi_xlsx")
    ap.add_argument("-o", "--out", default="FAMILY", help="output directory (default: FAMILY)")
    ap.add_argument("-j", "--jobs", type=int, default=6)
    args = ap.parse_args(argv)

    base = os.path.basename(os.path.normpath(args.xlsx_dir)).removesuffix("_xlsx")
    files = sorted(glob.glob(os.path.join(args.xlsx_dir, "*.xlsx")), key=part_no)
    if not files:
        sys.exit(f"no workbooks in {args.xlsx_dir}")
    with Pool(args.jobs) as pool:
        results = sorted(pool.imap_unordered(process, files, chunksize=4))

    book = Workbook(write_only=True)
    summ = book.create_sheet("Summary")
    fams_ws = book.create_sheet("Families")
    mem_ws = book.create_sheet("Members")
    summ.append(["Part No", "Voters", "Families", "Voters In Families", "% In Families",
                 "Largest Family"])
    fams_ws.append(["Part No", "Family ID", "Size", "House No(s)",
                    "Family Members (Name, EPIC, Serial No, Sex, Age)"])
    mem_ws.append(["Part No", "Family ID", "Serial No", "EPIC", "Name", "Age", "Gender",
                   "Relation Type", "Relation Name", "House No", "Inferred Links"])
    js, tot_v, tot_f, tot_in = {}, 0, 0, 0
    for part, voters, fams, links in results:
        live = sum(1 for v in voters if not v["deleted"])
        inf = sum(len(f) for f in fams)
        tot_v, tot_f, tot_in = tot_v + live, tot_f + len(fams), tot_in + inf
        summ.append([part, live, len(fams), inf, round(100 * inf / live, 1) if live else None,
                     max((len(f) for f in fams), default=0)])
        js[part] = []
        for k, f in enumerate(fams, 1):
            fid = f"{part}-F{k}"
            houses = sorted({voters[i]["house"] for i in f if house_key(voters[i]["house"])})
            fams_ws.append([part, fid, len(f), ", ".join(houses), "; ".join(
                f"{voters[i]['name']} ({voters[i]['epic']}, {voters[i]['sl']}, "
                f"{voters[i]['gender'][:1] or '?'}, {voters[i]['age'] if voters[i]['age'] is not None else '?'})"
                for i in f)])
            for i in f:
                v = voters[i]
                mem_ws.append([part, fid, v["sl"], v["epic"], v["name"], v["age"], v["gender"],
                               v["rtype"].title(), v["rname"], v["house"],
                               "; ".join(f"{r} of {s}" for r, s in links.get(i, []))])
            js[part].append([k, [voters[i]["sl"] for i in f],
                             [[[r, s] for r, s in links.get(i, [])] for i in f]])
    summ.append(["TOTAL", tot_v, tot_f, tot_in, round(100 * tot_in / tot_v, 1) if tot_v else None, None])
    os.makedirs(args.out, exist_ok=True)
    out = os.path.join(args.out, f"{base}_Families.xlsx")
    book.save(out)
    with open(os.path.join(args.out, f"{base}_Families.json"), "w", encoding="utf-8") as fh:
        json.dump(js, fh, ensure_ascii=False, separators=(",", ":"))
    print(f"{out}: {len(results)} booths, {tot_v} voters, {tot_f} families, "
          f"{tot_in} voters in families ({100 * tot_in / max(1, tot_v):.1f}%)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
